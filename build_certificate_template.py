from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "output/editable_certificate_template.docx"
NAVY = "17365D"
GOLD = "C69C3C"
GRAY = "595959"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color, size="18"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)
        borders.append(element)


def add_centered(doc, text, size, color="000000", bold=False, space_before=0, space_after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Aptos Display"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    return p


doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = section.page_height, section.page_width
section.top_margin = Inches(0.42)
section.bottom_margin = Inches(0.42)
section.left_margin = Inches(0.52)
section.right_margin = Inches(0.52)

# Decorative double border, made from nested tables for stable Word/LibreOffice rendering.
outer = doc.add_table(rows=1, cols=1)
outer.autofit = False
outer.columns[0].width = Inches(9.95)
outer_cell = outer.cell(0, 0)
set_cell_shading(outer_cell, "FFFFFF")
set_cell_border(outer_cell, NAVY, "26")

inner = outer_cell.add_table(rows=1, cols=1, width=Inches(9.63))
inner_cell = inner.cell(0, 0)
set_cell_border(inner_cell, GOLD, "10")

# All replaceable values are bracketed placeholders so the file is immediately usable in Word.
p = inner_cell.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(14)
p.paragraph_format.space_after = Pt(4)
r = p.add_run("[ORGANIZATION NAME]")
r.bold = True
r.font.name = "Aptos Display"
r.font.size = Pt(14)
r.font.color.rgb = RGBColor.from_string(NAVY)

p = inner_cell.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(10)
r = p.add_run("CERTIFICATE OF COMPLETION")
r.bold = True
r.font.name = "Aptos Display"
r.font.size = Pt(31)
r.font.color.rgb = RGBColor.from_string(NAVY)

p = inner_cell.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
r = p.add_run("This certificate is proudly presented to")
r.font.name = "Aptos"
r.font.size = Pt(13)
r.font.color.rgb = RGBColor.from_string(GRAY)

p = inner_cell.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(9)
r = p.add_run("[RECIPIENT NAME]")
r.bold = True
r.font.name = "Georgia"
r.font.size = Pt(27)
r.font.color.rgb = RGBColor.from_string(GOLD)

p = inner_cell.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(5)
r = p.add_run("for successfully completing")
r.font.name = "Aptos"
r.font.size = Pt(13)
r.font.color.rgb = RGBColor.from_string(GRAY)

p = inner_cell.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(18)
r = p.add_run("[COURSE OR ACHIEVEMENT TITLE]")
r.bold = True
r.font.name = "Aptos Display"
r.font.size = Pt(18)
r.font.color.rgb = RGBColor.from_string(NAVY)

info = inner_cell.add_table(rows=2, cols=3, width=Inches(8.85))
info.autofit = False
widths = [2.7, 3.45, 2.7]
for row in info.rows:
    for idx, cell in enumerate(row.cells):
        cell.width = Inches(widths[idx])
        for par in cell.paragraphs:
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER

labels = ["DATE", "CERTIFICATE ID", "AUTHORIZED SIGNATURE"]
values = ["[DD MONTH YYYY]", "[CERTIFICATE ID]", "[SIGNATORY NAME]"]
for i in range(3):
    p = info.cell(0, i).paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(labels[i])
    r.bold = True
    r.font.name = "Aptos"
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p = info.cell(1, i).paragraphs[0]
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(values[i])
    r.font.name = "Aptos"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(GRAY)

p = inner_cell.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(1)
p.paragraph_format.space_after = Pt(11)
r = p.add_run("Replace all text in [BRACKETS] with your details before issuing this certificate.")
r.italic = True
r.font.name = "Aptos"
r.font.size = Pt(8)
r.font.color.rgb = RGBColor.from_string(GRAY)

doc.save(OUT)
print(OUT)
